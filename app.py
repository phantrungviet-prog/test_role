import html
import os
import re
import streamlit as st
from google import genai
from google.genai import types

# 1. Cấu hình trang Streamlit
st.set_page_config(page_title="PoC Trợ lý Pháp lý & Kiểm toán Hợp đồng AI", layout="wide")

# Giao diện màu "Midnight Executive" (navy / ice / vàng nhấn), màu nền theo .streamlit/config.toml
st.markdown("""
<style>
.app-header {
    background: #1E2761;
    border-bottom: 4px solid #F9E795;
    border-radius: 10px;
    padding: 22px 28px;
    margin-bottom: 18px;
}
.app-header h1 { color: #FFFFFF; font-size: 1.7rem; margin: 0; padding: 0; }
.app-header p  { color: #CADCFC; margin: 6px 0 0; font-size: 0.95rem; }
[data-testid="stChatMessage"] {
    background: #FFFFFF;
    border: 1px solid #CADCFC;
    border-radius: 10px;
}
.status { padding: 8px 12px; margin: 6px 0; border-radius: 6px; border-left: 4px solid; }
.status-red   { background: #FDF1F1; border-color: #B00020; }
.status-amber { background: #FCF6E3; border-color: #B8860B; }
.status-green { background: #EDF5EC; border-color: #2C5F2D; }
</style>
<div class="app-header">
  <h1>🏛️ Trợ lý AI Pháp lý Du lịch &amp; PCCC</h1>
  <p>Tra cứu quy định · Kiểm toán hợp đồng · Phát hiện rủi ro tuân thủ</p>
</div>
""", unsafe_allow_html=True)

_STATUS_CLASS = {"🔴": "status-red", "🟡": "status-amber", "🟢": "status-green"}

def render_reply(text):
    """Hiển thị câu trả lời; dòng có 🔴/🟡/🟢 được tô thành khung màu theo mức độ."""
    out = []
    for line in text.splitlines():
        cls = next((c for icon, c in _STATUS_CLASS.items() if icon in line), None)
        if cls:
            body = html.escape(line.strip().lstrip("-* ").strip())
            body = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", body)
            out.append(f'<div class="status {cls}">{body}</div>\n')
        else:
            out.append(line)
    st.markdown("\n".join(out), unsafe_allow_html=True)

# 2. Lấy API Key: ưu tiên secrets.toml / biến môi trường, nếu không có thì nhập ở sidebar
def _load_api_key():
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        key = None
    return key or os.environ.get("GEMINI_API_KEY")

api_key = _load_api_key() or st.sidebar.text_input("Nhập Google Gemini API Key:", type="password")
if not api_key:
    st.warning("Vui lòng nhập API Key để bắt đầu sử dụng.")
    st.stop()

# Khởi tạo Gemini Client
client = genai.Client(api_key=api_key)

# 3. Khu vực Upload tài liệu pháp luật và hợp đồng
st.sidebar.header("📁 Kho tài liệu & Hợp đồng đầu vào")
uploaded_files = st.sidebar.file_uploader(
    "Tải lên các file Luật (PDF) và Hợp đồng cần kiểm tra:",
    type=["pdf", "txt", "docx"],
    accept_multiple_files=True
)

# 4. Định nghĩa System Prompt (Rất quan trọng để ép AI tuân thủ)
SYSTEM_INSTRUCTION = """
[VAI TRÒ VÀ BẢN THỂ]
Bạn là "Trợ lý AI Pháp lý & Kiểm toán Tuân thủ chuyên sâu" dành riêng cho ngành Du lịch, Khách sạn, Lữ hành và An toàn Phòng cháy chữa cháy (PCCC) tại Việt Nam.
Mục tiêu của bạn là giúp doanh nghiệp tra cứu văn bản pháp luật chính xác 100% và tự động phát hiện rủi ro, lỗi vi phạm trong các hợp đồng/chứng từ được tải lên.

[NGUỒN DỮ LIỆU BẮT BUỘC & QUY TẮC CHỐNG ẢO GIÁC]
1. Bạn CHỈ ĐƯỢC PHÉP dựa vào nội dung các văn bản pháp luật và tệp tài liệu do người dùng đính kèm trong phiên làm việc để làm căn cứ pháp lý.
2. Tuyệt đối KHÔNG tự suy diễn, KHÔNG tự bịa đặt số hiệu điều luật hoặc nội dung không có trong tài liệu đính kèm.
3. Nếu tài liệu đính kèm không có thông tin để trả lời, phải trả lời rõ ràng: "Rất tiếc, thông tin này không có trong các văn bản/tài liệu được cung cấp."

[NHIỆM VỤ CHÍNH]

Nhiệm vụ 1: Tra cứu Pháp lý & Quy định Ngành
- Đọc, tổng hợp và giải đáp thắc mắc của người dùng về Luật Du lịch, Điều kiện kinh doanh lữ hành, Tiêu chuẩn xếp hạng cơ sở lưu trú và Quy chuẩn PCCC.
- Mọi câu trả lời BẮT BUỘC phải kèm trích dẫn chính xác định dạng: [Tên văn bản | Điều... | Khoản...].

Nhiệm vụ 2: Kiểm toán & Thẩm định Hợp đồng / Chứng từ (Document Auditing)
Khi người dùng tải lên hợp đồng (như hợp đồng thuê mặt bằng, hợp đồng lữ hành, cam kết PCCC...), hãy thực hiện kiểm toán chéo theo 4 bước và xuất báo cáo:
1. Bóc tách thực thể: Xác định bên A, bên B, phạm vi, giá trị, thời hạn.
2. Đối chiếu pháp lý: So sánh các điều khoản hợp đồng với quy định trong các file luật đính kèm.
3. Chấm điểm & Phân loại rủi ro: Đánh giá theo 3 mức độ:
   - 🟢 [ĐẠT]: Điều khoản hợp pháp, đầy đủ.
   - 🟡 [CẢNH BÁO RỦI RO]: Thiếu điều khoản bảo vệ quyền lợi, câu chữ mơ hồ hoặc mức phạt vượt quy định.
   - 🔴 [VI PHẠM PHÁP LUẬT]: Trái quy định bắt buộc, thiếu giấy phép con/cam kết PCCC bắt buộc.
4. Xuất Báo cáo Kiểm toán (Compliance Audit Report) cấu trúc rõ ràng.

[ĐỊNH DẠNG BÁO CÁO KIỂM TOÁN MẪU]
Khi phân tích hợp đồng, hãy trả về kết quả theo cấu trúc sau:

📋 BÁO CÁO KIỂM TOÁN TÍNH HỢP LỆ HỢP ĐỒNG
- Tên tài liệu phân tích: [Tên file hợp đồng]
- Đánh giá tổng thể: [ĐẠT / CẦN SỬA ĐỔI / VI PHẠM nghiêm trọng]

1. ĐIỂM TỰA PHÁP LÝ & ĐIỀU KIỆN BẮT BUỘC:
   - Status: [🟢/🟡/🔴] [Tóm tắt phát hiện]
   - Chi tiết: [Diễn giải vấn đề]
   - Căn cứ pháp lý: [Tên văn bản đính kèm - Điều... Khoản...]

2. ĐIỀU KHOẢN THƯƠNG MẠI & RỦI RO TRANH CHẤP:
   - Status: [🟢/🟡/🔴] [Tóm tắt phát hiện]
   - Chi tiết: [Diễn giải rủi ro]

3. ĐỀ XUẤT ĐIỀU CHỈNH TỪ AI:
   - [Nêu rõ câu chữ cần sửa hoặc điều khoản cần bổ sung]

[TÔNG GIỌNG]
Khách quan, chuẩn xác, ngôn ngữ pháp lý doanh nghiệp, ngắn gọn và trực diện.
"""

def extract_docx_text(file):
    """Trích văn bản (đoạn + bảng) từ file .docx."""
    import io
    import docx  # pip install python-docx

    document = docx.Document(io.BytesIO(file.getvalue()))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)

# 5. Giao diện Chat chính
if "messages" not in st.session_state:
    st.session_state.messages = []

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        if message["role"] == "assistant":
            render_reply(message["content"])
        else:
            st.markdown(message["content"])

user_query = st.chat_input("Nhập câu hỏi pháp lý hoặc yêu cầu kiểm tra hợp đồng...")

if user_query:
    st.session_state.messages.append({"role": "user", "content": user_query})
    with st.chat_message("user"):
        st.markdown(user_query)

    with st.chat_message("assistant"):
        with st.spinner("AI đang đọc tài liệu và phân tích..."):
            try:
                # Lịch sử hội thoại (trừ câu hỏi vừa nhập) để AI nhớ các lượt trước
                contents = [
                    types.Content(
                        role="user" if m["role"] == "user" else "model",
                        parts=[types.Part.from_text(text=m["content"])],
                    )
                    for m in st.session_state.messages[:-1]
                ]

                # Lượt hiện tại: tài liệu đính kèm + câu hỏi
                current_parts = []
                if uploaded_files:
                    for file in uploaded_files:
                        # Ghi tên file để AI trích dẫn đúng "Tên tài liệu"
                        current_parts.append(types.Part.from_text(text=f"[Tài liệu: {file.name}]"))
                        if file.name.lower().endswith(".docx"):
                            # Gemini không đọc trực tiếp .docx -> chuyển sang văn bản
                            current_parts.append(types.Part.from_text(text=extract_docx_text(file)))
                        else:
                            current_parts.append(
                                types.Part.from_bytes(
                                    data=file.getvalue(),
                                    mime_type=file.type if file.type else "application/pdf"
                                )
                            )
                current_parts.append(types.Part.from_text(text=user_query))
                contents.append(types.Content(role="user", parts=current_parts))

                # Gọi model Gemini Flash (tối ưu tốc độ và context lớn)
                response = client.models.generate_content(
                    model='gemini-3.8-flash',
                    contents=contents,
                    config=types.GenerateContentConfig(
                        system_instruction=SYSTEM_INSTRUCTION,
                        temperature=0.1, # Giảm sáng tạo để đảm bảo tính chính xác pháp lý
                    )
                )


                
                bot_reply = response.text
                render_reply(bot_reply)
                st.session_state.messages.append({"role": "assistant", "content": bot_reply})
                
            except Exception as e:
                st.error(f"Đã xảy ra lỗi khi gọi Gemini API: {e}")